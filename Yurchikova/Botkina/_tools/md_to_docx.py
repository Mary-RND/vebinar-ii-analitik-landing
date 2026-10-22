# -*- coding: utf-8 -*-
"""Convert the project's Markdown deliverables to styled DOCX files."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ACCENT = RGBColor(0x1F, 0x3A, 0x5F)
MUTED = RGBColor(0x55, 0x5F, 0x6D)


def set_run_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:cs"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_page_number_footer(section):
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Стр. ")
    set_run_font(run, size=8, color=MUTED)
    run2 = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run2._r.append(fld_begin)
    run2._r.append(instr)
    run2._r.append(fld_end)
    set_run_font(run2, size=8, color=MUTED)


def add_header(section, text):
    p = section.header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    set_run_font(run, size=8, color=MUTED)


def start_landscape(doc):
    """Новый раздел A4 с альбомной ориентацией — для широких таблиц."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.LANDSCAPE
    s.page_width = Cm(29.7)
    s.page_height = Cm(21.0)
    s.top_margin = Cm(1.7)
    s.bottom_margin = Cm(1.6)
    s.left_margin = Cm(1.5)
    s.right_margin = Cm(1.5)
    return s


def end_landscape(doc):
    """Возврат к книжной ориентации после широкой таблицы."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    s.orientation = WD_ORIENT.PORTRAIT
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.top_margin = Cm(1.9)
    s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.2)
    s.right_margin = Cm(1.6)
    return s


def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = OxmlElement("w:tblCellMar")
    for name, value in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = OxmlElement(f"w:{name}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        tbl_cell_mar.append(el)
    tbl_pr.append(tbl_cell_mar)


def mark_header_row(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_bottom_border(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "B8C2CC")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*|`[^`]+?`)")


def add_inline(paragraph, text, size=None, color=None, base_bold=False):
    for token in INLINE_RE.split(text):
        if not token:
            continue
        bold = base_bold
        italic = False
        mono = False
        if token.startswith("**") and token.endswith("**"):
            token = token[2:-2]
            bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            token = token[1:-1]
            italic = True
        elif token.startswith("`") and token.endswith("`"):
            token = token[1:-1]
            mono = True
        run = paragraph.add_run(token)
        set_run_font(
            run,
            name="Consolas" if mono else "Times New Roman",
            size=size,
            bold=bold,
            italic=italic,
            color=color,
        )


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def add_table(doc, rows):
    if not rows:
        return
    header = rows[0]
    data = rows[1:]
    header_empty = all(not cell for cell in header)
    if header_empty:
        table = doc.add_table(rows=len(data), cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_cell_margins(table)
        for r, row_data in enumerate(data):
            for i, cell_text in enumerate(row_data[: len(header)]):
                cell = table.rows[r].cells[i]
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline(p, cell_text, size=8.0)
    else:
        table = doc.add_table(rows=1, cols=len(header))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        set_cell_margins(table)
        hdr = table.rows[0]
        mark_header_row(hdr)
        for i, cell_text in enumerate(header):
            cell = hdr.cells[i]
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_inline(p, cell_text, size=8.0, color=ACCENT, base_bold=True)
            for run in p.runs:
                run.bold = True
        for row_data in data:
            row = table.add_row()
            for i, cell_text in enumerate(row_data[: len(header)]):
                cell = row.cells[i]
                cell.text = ""
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                add_inline(p, cell_text, size=8.0)
    # Keep table rows from splitting across pages.
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        cant_split.set(qn("w:val"), "true")
        tr_pr.append(cant_split)
    doc.add_paragraph()


def add_rule(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    add_bottom_border(p)


def convert(md_path: Path, docx_path: Path, header_text: str, base_dir: Path | None = None):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()
    img_base = base_dir or md_path.parent

    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.9)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(1.6)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (("Heading 1", 15, ACCENT), ("Heading 2", 12.5, ACCENT), ("Heading 3", 11, ACCENT)):
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:cs"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(10)
        st.paragraph_format.space_after = Pt(4)
        st.paragraph_format.keep_with_next = True

    add_header(section, header_text)
    add_page_number_footer(section)

    first_h1 = True
    first_h2 = True
    first_rule = True
    i = 0

    while i < len(lines):
        raw = lines[i]
        line = raw.rstrip()
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        # Do not include the internal service block in client deliverables.
        if stripped.startswith("### Служебный блок"):
            break

        # Explicit page break
        if stripped == "[[pagebreak]]":
            doc.add_page_break()
            i += 1
            continue

        # Markdown tables
        if stripped.startswith("|") and i + 1 < len(lines) and re.match(r"^\s*\|[\s:\-|]+\|\s*$", lines[i + 1]):
            table_rows = [split_table_row(stripped)]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_rows.append(split_table_row(lines[i].strip()))
                i += 1
            wide = len(table_rows[0]) >= 10
            if wide:
                start_landscape(doc)
            add_table(doc, table_rows)
            if wide:
                end_landscape(doc)
            continue

        # Horizontal rule / first section break
        if stripped == "---":
            if first_rule:
                doc.add_page_break()
                first_rule = False
            else:
                add_rule(doc)
            i += 1
            continue

        # Headings
        if stripped.startswith("# "):
            text_line = stripped[2:].strip()
            # If the next H2 is a numbered section, this H1 is the document title.
            next_h2 = ""
            for lookahead in lines[i + 1 :]:
                if lookahead.strip().startswith("## "):
                    next_h2 = lookahead.strip()[3:].strip()
                    break
                if lookahead.strip() and not lookahead.strip().startswith("**"):
                    break
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            if re.match(r"^\d+\.", next_h2):
                run = p.add_run(text_line)
                set_run_font(run, size=21, bold=True, color=ACCENT)
            else:
                run = p.add_run(text_line.upper())
                set_run_font(run, size=12, bold=True, color=MUTED)
            first_h1 = False
            i += 1
            continue

        if stripped.startswith("## "):
            text_line = stripped[3:].strip()
            if first_h2 and not re.match(r"^\d+\.", text_line):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run(text_line)
                set_run_font(run, size=21, bold=True, color=ACCENT)
                first_h2 = False
            else:
                doc.add_heading(text_line, level=1)
                first_h2 = False
            i += 1
            continue

        if stripped.startswith("### "):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith("#### "):
            doc.add_heading(stripped[5:].strip(), level=3)
            i += 1
            continue

        # Bullets
        bullet_match = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if bullet_match:
            indent = len(bullet_match.group(1))
            text_line = bullet_match.group(2)
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            add_inline(p, text_line, size=10.5)
            i += 1
            continue

        # Numbered items: keep literal numbering, avoid Word list continuation issues
        num_match = re.match(r"^(\d+)\.\s+(.*)$", stripped)
        if num_match:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.65)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(f"{num_match.group(1)}.\t")
            set_run_font(run, size=10.5, bold=False)
            add_inline(p, num_match.group(2), size=10.5)
            i += 1
            continue

        # Рисунок (ТЗ § 7.2)
        if stripped.startswith("!["):
            m = re.match(r"!\[([^\]]*)\]\(([^)\s]+)\)", stripped)
            if m:
                alt, src = m.group(1), m.group(2)
                img_path = (img_base / src)
                if img_path.exists():
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_before = Pt(4)
                    p.paragraph_format.space_after = Pt(2)
                    p.add_run().add_picture(str(img_path), width=Cm(16.8))
                    cap = doc.add_paragraph()
                    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cap.paragraph_format.space_after = Pt(8)
                    run = cap.add_run(alt)
                    set_run_font(run, size=8.5, italic=True, color=RGBColor(0x45, 0x50, 0x5D))
                else:
                    p = doc.add_paragraph()
                    add_inline(p, f"[Рисунок не найден: {src}]", size=9)
                i += 1
                continue

        # Regular paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_inline(p, stripped, size=10.5)
        i += 1

    doc.core_properties.title = header_text
    doc.core_properties.author = "Аналитический проект"
    doc.core_properties.language = "ru-RU"
    doc.save(docx_path)
    print(f"Saved {docx_path} ({docx_path.stat().st_size} bytes)")


if __name__ == "__main__":
    base = Path(__file__).resolve().parent.parent
    convert(
        base / "TZ_Botkin_2025.md",
        base / "TZ_Botkin_2025.docx",
        "Техническое задание — анализ закупок МНКЦ им. С.П. Боткина за 2025 год",
    )
    convert(
        base / "KP_Botkin_2025.md",
        base / "KP_Botkin_2025.docx",
        "Коммерческое предложение — анализ закупок МНКЦ им. С.П. Боткина за 2025 год",
    )
