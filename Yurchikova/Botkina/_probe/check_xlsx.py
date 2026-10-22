from pathlib import Path

from openpyxl import load_workbook

p = next(Path("output").glob("*.xlsx"))
wb = load_workbook(p, read_only=False)
print("file:", p.name, "sheets:", len(wb.sheetnames))
for ws in wb.worksheets:
    print(f"{ws.title!r}: dims={ws.dimensions} freeze={ws.freeze_panes} autofilter={ws.auto_filter.ref} merged={len(ws.merged_cells.ranges)}")
