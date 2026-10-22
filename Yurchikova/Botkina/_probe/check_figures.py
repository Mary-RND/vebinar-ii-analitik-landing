"""Проверка наличия рисунков в PDF и DOCX отчёта (ТЗ § 7.2)."""
from pathlib import Path

import fitz
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
PDF = ROOT / "output" / "REPORT_Botkin_2025.pdf"
DOCX = ROOT / "output" / "REPORT_Botkin_2025.docx"

doc = fitz.open(PDF)
pages = []
for i, page in enumerate(doc, 1):
    imgs = page.get_images(full=True)
    if imgs:
        rects = [page.get_image_rects(x[0]) for x in imgs]
        pages.append((i, len(imgs)))
print(f"PDF: {len(doc)} страниц, страниц с рисунками — {len(pages)}")
total = 0
for p, n in pages:
    print(f"  стр. {p}: рисунков {n}")
    total += n
print("всего изображений в PDF:", total)

# подписи «Рисунок N —»
caps = 0
for i, page in enumerate(doc, 1):
    txt = page.get_text()
    for line in txt.splitlines():
        if line.strip().startswith("Рисунок"):
            caps += 1
            print(f"  подпись стр. {i}: {line.strip()[:110]}")
print("подписей «Рисунок»:", caps)

d = Document(DOCX)
print("\nDOCX: рисунков (inline_shapes):", len(d.inline_shapes))
print("DOCX: таблиц:", len(d.tables))
caps_docx = 0
for p in d.paragraphs:
    if p.text.strip().startswith("Рисунок"):
        caps_docx += 1
print("DOCX: подписей «Рисунок»:", caps_docx)
