"""Проверка реестра позиций: состав полей (ТЗ § 5.4), заполненность, ссылки, CSV."""
import csv
import json
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "output" / "ММНКЦ_Боткина_2025_аналитика.xlsx"
CSV = ROOT / "output" / "Реестр_позиций_Боткина_2025.csv"

REQUIRED = {
    "position_id": ["position_id"],
    "исходное наименование": ["Исходное наименование"],
    "нормализованное наименование": ["Нормализованное наименование"],
    "ОКПД2": ["ОКПД2"],
    "КТРУ": ["КТРУ"],
    "ед. изм. (исходная/норм.)": ["Ед. изм. (исходная)", "Ед. изм. (норм.)"],
    "количество": ["Количество"],
    "цена за ед.": ["Цена за единицу, ₽"],
    "сумма позиции": ["Сумма, ₽"],
    "НДС": ["НДС"],
    "страна": ["Страна"],
    "производитель": ["Производитель"],
    "бренд": ["Бренд / товарный знак"],
    "модель": ["Модель"],
    "рег. удостоверение": ["Регистрационное удостоверение"],
    "поставщик": ["Поставщик"],
    "ИНН": ["ИНН поставщика"],
    "КПП": ["КПП поставщика"],
    "регион": ["Регион поставщика"],
    "scope": ["Scope"],
    "причина исключения": ["Причина классификации"],
    "L1": ["Вид (L1)"],
    "L2": ["Товарная группа (L2)"],
    "L3": ["Номенклатурная позиция (L3)"],
    "смешанный лот": ["Смешанный лот"],
    "confidence": ["Класс (факт/оценка)"],
    "комментарий": ["Комментарий"],
    "источник атрибуции": ["Источник атрибуции страны", "Источник атрибуции позиции"],
    "признак дубля": ["Признак дубля"],
    "ссылка на источник": ["Ссылка на источник"],
    "дата выгрузки": ["Дата выгрузки"],
}

wb = openpyxl.load_workbook(XLSX, read_only=True, data_only=True)
print("листов:", len(wb.sheetnames))
for n in wb.sheetnames:
    print("  ", n)

ws = wb["08_Реестр_позиций"]
rows = ws.iter_rows(values_only=True)
header = [str(h) for h in next(rows)]
data = [r for r in rows]
# последние строки — примечание и легенда
data = [r for r in data if r and isinstance(r[0], str) and r[0].startswith("P")]
print(f"\nреестр: столбцов {len(header)}, строк {len(data)}")
print("столбцы:")
for i, h in enumerate(header, 1):
    print(f"  {i:>2}. {h}")

print("\n=== соответствие полям ТЗ § 5.4 ===")
missing = []
for label, cols in REQUIRED.items():
    hit = [c for c in cols if c in header]
    if not hit:
        missing.append(label)
        print(f"  ОТСУТСТВУЕТ: {label}")
    else:
        filled = sum(1 for r in data if any(
            (str(r[header.index(c)]).strip() not in ("", "None", "—"))
            for c in hit if c in header))
        print(f"  {label:<34} → {', '.join(hit):<45} заполнено {filled:>5} "
              f"({filled / len(data) * 100:5.1f} %)")

# ссылка на источник
idx = header.index("Ссылка на источник")
links = [str(r[idx]) for r in data if r[idx]]
url = sum(1 for l in links if l.startswith("http"))
print(f"\nссылка на источник: строк со значением {len(links)} ({len(links) / len(data) * 100:.1f} %), "
      f"из них URL — {url} ({url / len(data) * 100:.1f} %)")

# объединённые ячейки
wb2 = openpyxl.load_workbook(XLSX)
merged = sum(len(w.merged_cells.ranges) for w in wb2.worksheets)
print("объединённых ячеек в книге:", merged)
print("скрытых строк/столбцов:",
      sum(1 for w in wb2.worksheets for d in w.row_dimensions.values() if d.hidden),
      sum(1 for w in wb2.worksheets for d in w.column_dimensions.values() if d.hidden))

print("\n=== CSV ===")
raw = CSV.read_bytes()
print("размер:", len(raw), "байт; BOM:", raw[:3] == b"\xef\xbb\xbf")
with CSV.open(encoding="utf-8-sig", newline="") as fh:
    rdr = csv.reader(fh, delimiter=";")
    head = next(rdr)
    n = sum(1 for _ in rdr)
print("строк (без шапки):", n, "| столбцов:", len(head))
print("разделитель ';':", head[1][:20])
