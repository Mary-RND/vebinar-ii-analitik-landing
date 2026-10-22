from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK

ROOT = Path(__file__).resolve().parents[1]
d = Document(ROOT / "output" / "REPORT_Botkin_2025.docx")
print("paragraphs:", len(d.paragraphs), "tables:", len(d.tables))
print("sections:", len(d.sections))
for i, s in enumerate(d.sections):
    print(f"  section {i}: page {s.page_width.mm:.0f}x{s.page_height.mm:.0f} мм, orient={s.orientation}")
    for name, part in (("header", s.header), ("footer", s.footer)):
        txt = " | ".join(p.text for p in part.paragraphs if p.text.strip())
        xml = part._element.xml if hasattr(part, "_element") else ""
        has_field = "PAGE" in xml
        print(f"    {name}: {txt!r} page_field={has_field}")

brk = 0
heads = []
for p in d.paragraphs:
    xml = p._p.xml
    if 'w:br' in xml and 'type="page"' in xml:
        brk += 1
    if p.style.name.startswith("Heading"):
        heads.append((p.style.name, p.text))
print("page breaks:", brk)
print("headings:", len(heads))
for h in heads[:40]:
    print("   ", h)
